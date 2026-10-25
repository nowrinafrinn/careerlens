"""Builds small PDF and DOCX files in memory for tests, so no binary fixtures
(and no real person's CV) need to live in the repository."""

import io

from docx import Document


def _pdf_escape(text: str) -> bytes:
    raw = text.encode("cp1252")  # WinAnsiEncoding, so "•" and "–" survive
    out = bytearray()
    for byte in raw:
        if byte in b"()\\":
            out += b"\\" + bytes([byte])
        elif byte > 126:
            out += b"\\%03o" % byte
        else:
            out.append(byte)
    return bytes(out)


def make_pdf(lines: list[str]) -> bytes:
    """A one-page, text-based PDF with one line of Helvetica text per item."""
    content = bytearray(b"BT /F1 10 Tf 13 TL 50 760 Td\n")
    for line in lines:
        content += b"(" + _pdf_escape(line) + b") Tj T*\n"
    content += b"ET"

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + bytes(content) + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref_at)
    return bytes(out)


def make_docx(blocks: list[tuple[str, str | None]], table: list[list[str]] | None = None) -> bytes:
    """blocks: (text, style) pairs, e.g. ("Education", "Heading 1") or ("BSc", None)."""
    document = Document()
    for text, style in blocks:
        document.add_paragraph(text, style=style) if style else document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


SAMPLE_CV_LINES = [
    "Rafi Ahmed",
    "rafi@example.com | Dhaka | github.com/rafi-example",
    "Summary",
    "CSE graduate interested in backend development.",
    "Education",
    "BSc in Computer Science and Engineering, Example University, 2021-2025",
    "Technical Skills",
    "Languages: Python, TypeScript, SQL",
    "Tools: Docker, Git, PostgreSQL",
    "Experience",
    "Software Engineering Intern, Example Tech Ltd, Jun-Aug 2024",
    "• Built REST endpoints with FastAPI for an internal reporting tool,",
    "reducing manual report preparation.",
    "• Wrote pytest tests for the data import module.",
    "Projects",
    "CareerLens - job application tracker",
    "• Designed a PostgreSQL schema with Alembic migrations.",
]
